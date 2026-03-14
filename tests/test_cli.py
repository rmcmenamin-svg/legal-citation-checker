"""Tests for the CLI interface."""

import sys
from pathlib import Path
from typing import Any, Dict, Tuple
from unittest.mock import patch

import pytest

from legal_citation_checker.cli import main


class TestCLIArgs:
    def test_version_flag(self, capsys: pytest.CaptureFixture[str]) -> None:
        with pytest.raises(SystemExit) as exc_info:
            main(["--version"])
        assert exc_info.value.code == 0
        captured = capsys.readouterr()
        assert "0.1.0" in captured.out

    def test_missing_file_returns_1(self, capsys: pytest.CaptureFixture[str]) -> None:
        result = main(["nonexistent.docx"])
        assert result == 1
        captured = capsys.readouterr()
        assert "not found" in captured.err.lower()

    def test_unsupported_file_type_returns_1(self, tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
        txt_file = tmp_path / "test.txt"
        txt_file.write_text("hello")
        result = main([str(txt_file)])
        assert result == 1
        captured = capsys.readouterr()
        assert "supported" in captured.err.lower()

    def test_timeout_option_accepted(self, tmp_path: Path) -> None:
        try:
            from docx import Document  # type: ignore
        except ImportError:
            pytest.skip("python-docx not installed")

        doc_path = tmp_path / "test.docx"
        doc = Document()
        doc.add_paragraph("No citations here.")
        doc.save(str(doc_path))

        result = main([str(doc_path), "--timeout", "15"])
        assert result == 0

    def test_no_cache_option_accepted(self, tmp_path: Path) -> None:
        try:
            from docx import Document  # type: ignore
        except ImportError:
            pytest.skip("python-docx not installed")

        doc_path = tmp_path / "test.docx"
        doc = Document()
        doc.add_paragraph("No citations here.")
        doc.save(str(doc_path))

        result = main([str(doc_path), "--no-cache"])
        assert result == 0

    def test_json_format_output(self, tmp_path: Path) -> None:
        try:
            from docx import Document  # type: ignore
        except ImportError:
            pytest.skip("python-docx not installed")

        doc_path = tmp_path / "test.docx"
        doc = Document()
        doc.add_paragraph("No citations here.")
        doc.save(str(doc_path))

        output_path = tmp_path / "report.json"
        result = main([str(doc_path), "-f", "json", "-o", str(output_path)])
        assert result == 0
        assert output_path.exists()
        import json
        data = json.loads(output_path.read_text())
        assert "summary" in data
