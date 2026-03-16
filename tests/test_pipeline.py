"""Tests for the citation checker pipeline."""

import json
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pytest

from legal_citation_checker.pipeline import (
    AuditReport,
    CitationAudit,
    CitationChecker,
    ExtractedCitation,
    SearchAttempt,
    VerificationDecision,
    _SKIP_CITATION_TYPES,
    _WESTLAW_PATTERN,
    _extract_name_tokens,
    _name_token_overlap,
)


# ---------------------------------------------------------------------------
# Fixtures & helpers
# ---------------------------------------------------------------------------

def make_citation(
    index: int = 1,
    raw: str = "347 U.S. 483",
    normalized: str = "347 U.S. 483",
    citation_type: str = "FullCaseCitation",
    metadata: Optional[Dict[str, Any]] = None,
) -> ExtractedCitation:
    return ExtractedCitation(
        index=index,
        raw_citation=raw,
        normalized_citation=normalized,
        citation_type=citation_type,
        context=f"...see {raw}...",
        paragraph_index=1,
        metadata=metadata or {"year": "1954", "plaintiff": "Brown", "defendant": "Board of Education"},
        bluebook_normalized=False,
    )


def make_checker(**kwargs: Any) -> CitationChecker:
    return CitationChecker(verbose=False, max_workers=1, **kwargs)


# ---------------------------------------------------------------------------
# Unit tests: Bluebook normalization
# ---------------------------------------------------------------------------

class TestNormalizeBluebook:
    def test_normalizes_whitespace(self) -> None:
        checker = make_checker()
        result, changed = checker._normalize_bluebook("347  U.S.   483", {})
        assert "  " not in result

    def test_normalizes_section_symbol(self) -> None:
        checker = make_checker()
        result, _ = checker._normalize_bluebook("Section 1983", {})
        assert "§ 1983" in result

    def test_normalizes_v_separator(self) -> None:
        checker = make_checker()
        result, _ = checker._normalize_bluebook("Brown v Board", {})
        assert "v." in result

    def test_strips_trailing_semicolons(self) -> None:
        checker = make_checker()
        result, _ = checker._normalize_bluebook("347 U.S. 483;", {})
        assert not result.endswith(";")

    def test_preserves_already_correct_citation(self) -> None:
        checker = make_checker()
        result, changed = checker._normalize_bluebook("347 U.S. 483", {})
        assert result == "347 U.S. 483"
        assert not changed


# ---------------------------------------------------------------------------
# Unit tests: Statute/regulation detection
# ---------------------------------------------------------------------------

class TestStatuteDetection:
    def test_usc_detected(self) -> None:
        assert CitationChecker._is_statute_or_regulation("42 U.S.C. § 1983")

    def test_cfr_detected(self) -> None:
        assert CitationChecker._is_statute_or_regulation("29 C.F.R. Part 1926")

    def test_pub_law_detected(self) -> None:
        assert CitationChecker._is_statute_or_regulation("Pub. L. No. 111-148")

    def test_state_statute_detected(self) -> None:
        assert CitationChecker._is_statute_or_regulation("Cal. Code § 452")

    def test_case_citation_not_detected(self) -> None:
        assert not CitationChecker._is_statute_or_regulation("347 U.S. 483")

    def test_case_name_not_detected(self) -> None:
        assert not CitationChecker._is_statute_or_regulation("Brown v. Board of Education")


# ---------------------------------------------------------------------------
# Unit tests: Pincite stripping
# ---------------------------------------------------------------------------

class TestPinciteStripping:
    def test_strips_pincite(self) -> None:
        assert CitationChecker._strip_pincite("347 U.S. 483, 495") == "347 U.S. 483"

    def test_strips_pincite_range(self) -> None:
        assert CitationChecker._strip_pincite("347 U.S. 483, 495-497") == "347 U.S. 483"

    def test_no_pincite_unchanged(self) -> None:
        assert CitationChecker._strip_pincite("347 U.S. 483") == "347 U.S. 483"


# ---------------------------------------------------------------------------
# Unit tests: Westlaw detection
# ---------------------------------------------------------------------------

class TestWestlawDetection:
    def test_wl_citation_detected(self) -> None:
        assert _WESTLAW_PATTERN.match("2025 WL 2192378")

    def test_case_citation_not_detected(self) -> None:
        assert not _WESTLAW_PATTERN.match("347 U.S. 483")


# ---------------------------------------------------------------------------
# Unit tests: Citation type filtering
# ---------------------------------------------------------------------------

class TestCitationTypeFiltering:
    def test_skip_types_defined(self) -> None:
        for t in ("IdCitation", "SupraCitation", "ShortCaseCitation", "NonopinionCitation", "UnknownCitation"):
            assert t in _SKIP_CITATION_TYPES

    def test_full_case_not_skipped(self) -> None:
        assert "FullCaseCitation" not in _SKIP_CITATION_TYPES


# ---------------------------------------------------------------------------
# Unit tests: Cache key
# ---------------------------------------------------------------------------

class TestCacheKey:
    def test_canonical_text_lowered(self) -> None:
        checker = make_checker()
        assert checker._canonical_text("Brown V. BOARD") == "brown v board"

    def test_canonical_text_strips_punctuation(self) -> None:
        checker = make_checker()
        assert checker._canonical_text("347 U.S. 483") == "347 u s 483"

    def test_cache_key_deduplicates(self) -> None:
        checker = make_checker()
        assert checker._cache_key("347 U.S. 483") == checker._cache_key("347 U.S.  483")


# ---------------------------------------------------------------------------
# Unit tests: Reporter triplet
# ---------------------------------------------------------------------------

class TestReporterTriplet:
    def test_extracts_triplet(self) -> None:
        checker = make_checker()
        result = checker._reporter_triplet("347 U.S. 483")
        assert result is not None
        vol, reporter, page = result
        assert vol == "347"
        assert page == "483"

    def test_no_triplet_in_text(self) -> None:
        checker = make_checker()
        assert checker._reporter_triplet("some random text") is None

    def test_triplet_match_works(self) -> None:
        checker = make_checker()
        target = checker._reporter_triplet("347 U.S. 483")
        assert target is not None
        assert checker._triplet_match("347 U.S. 483", target)
        assert not checker._triplet_match("348 U.S. 483", target)


# ---------------------------------------------------------------------------
# Unit tests: Raw reporter extraction
# ---------------------------------------------------------------------------

class TestRawReporter:
    def test_extracts_reporter(self) -> None:
        assert CitationChecker._raw_reporter_from_citation("56 Cal. 2d 407") == "Cal. 2d"

    def test_extracts_reporter_with_app(self) -> None:
        assert CitationChecker._raw_reporter_from_citation("101 Cal. App. 5th 902") == "Cal. App. 5th"


# ---------------------------------------------------------------------------
# Unit tests: Needs Review vs Hallucination
# ---------------------------------------------------------------------------

class TestVerificationDecisions:
    def test_all_network_errors_with_zero_results_yields_hallucination(self) -> None:
        """When all sources error but returned 0 results, flag as hallucination."""
        checker = make_checker()
        citation = make_citation()

        def error_http_get(*args: Any, **kwargs: Any) -> Tuple[None, None, str]:
            return None, None, "Connection timeout"

        checker._verifier._http_get_json = error_http_get  # type: ignore[assignment]
        # Disable session so only CourtListener attempts are made
        checker._verifier._session = None  # type: ignore[assignment]

        decision = checker._verify_citation(citation)
        # Multiple error attempts with 0 results → hallucination
        assert decision.status == "Potential Hallucination"

    def test_all_network_errors_many_attempts_yields_hallucination(self) -> None:
        """When many sources are tried and all error with 0 results, flag hallucination."""
        checker = make_checker()
        citation = make_citation()

        def error_http_get(*args: Any, **kwargs: Any) -> Tuple[None, None, str]:
            return None, None, "Connection timeout"

        checker._verifier._http_get_json = error_http_get  # type: ignore[assignment]

        decision = checker._verify_citation(citation)
        # With enough error-but-empty attempts, we flag as potential hallucination
        assert decision.status == "Potential Hallucination"

    def test_real_not_found_yields_hallucination(self) -> None:
        checker = make_checker()
        citation = make_citation()

        def empty_http_get(*args: Any, **kwargs: Any) -> Tuple[Dict[str, Any], str, None]:
            return {"count": 0, "results": []}, "http://example.com", None

        checker._verifier._http_get_json = empty_http_get  # type: ignore[assignment]

        decision = checker._verify_citation(citation)
        assert decision.status == "Potential Hallucination"
        assert decision.confidence > 0

    def test_verified_on_courtlistener_match(self) -> None:
        checker = make_checker()
        citation = make_citation()

        def mock_http_get(url: str, **kwargs: Any) -> Tuple[Dict[str, Any], str, None]:
            return {
                "count": 1,
                "results": [{"citation": ["347 U.S. 483"], "absolute_url": "/opinion/123/"}],
            }, url, None

        checker._verifier._http_get_json = mock_http_get  # type: ignore[assignment]
        decision = checker._verify_citation(citation)
        assert decision.status.startswith("Verified")
        assert decision.confidence >= 85

    def test_westlaw_citation_no_party_needs_review(self) -> None:
        """WL citation with no party names falls back to Needs Review."""
        checker = make_checker()
        citation = make_citation(
            raw="2025 WL 2192378",
            normalized="2025 WL 2192378",
            metadata={"year": "2025"},
        )
        decision = checker._verify_citation(citation)
        assert decision.status == "Needs Review"
        assert "WL" in decision.evidence

    def test_westlaw_citation_with_parties_tries_verification(self) -> None:
        """WL citation with party names attempts party-name verification."""
        checker = make_checker()
        citation = make_citation(
            raw="2024 WL 2208099",
            normalized="2024 WL 2208099",
            metadata={"year": "2024", "plaintiff": "Franklyn", "defendant": "Daubert"},
        )

        # Mock CourtListener to return a matching case
        def mock_http_get(url: str, **kwargs: Any) -> Tuple[Dict[str, Any], str, None]:
            return {
                "count": 1,
                "results": [{
                    "caseName": "Franklyn v. Daubert",
                    "case_name": "Franklyn v. Daubert",
                    "dateFiled": "2024-05-15",
                    "absolute_url": "/opinion/123/franklyn-v-daubert/",
                    "cluster_id": 123,
                }],
            }, "http://example.com", None

        checker._verifier._http_get_json = mock_http_get  # type: ignore[assignment]
        decision = checker._verify_citation(citation)
        assert decision.status == "Verified (case exists)"
        assert decision.confidence > 0
        assert "WL" in decision.evidence

    def test_lexis_citation_with_parties_tries_verification(self) -> None:
        """LEXIS citation with party names attempts party-name verification."""
        checker = make_checker()
        citation = make_citation(
            raw="2024 U.S. App. LEXIS 12345",
            normalized="2024 U.S. App. LEXIS 12345",
            metadata={"year": "2024", "plaintiff": "Smith", "defendant": "Jones"},
        )

        def mock_http_get(url: str, **kwargs: Any) -> Tuple[Dict[str, Any], str, None]:
            return {
                "count": 1,
                "results": [{
                    "caseName": "Smith v. Jones",
                    "case_name": "Smith v. Jones",
                    "dateFiled": "2024-03-01",
                    "absolute_url": "/opinion/456/smith-v-jones/",
                    "cluster_id": 456,
                }],
            }, "http://example.com", None

        checker._verifier._http_get_json = mock_http_get  # type: ignore[assignment]
        decision = checker._verify_citation(citation)
        assert decision.status == "Verified (case exists)"
        assert "LEXIS" in decision.evidence


# ---------------------------------------------------------------------------
# Unit tests: Audit report
# ---------------------------------------------------------------------------

class TestAuditReport:
    def _make_report(self, statuses: List[str]) -> AuditReport:
        citations = []
        for i, status in enumerate(statuses, 1):
            citations.append(
                CitationAudit(
                    index=i,
                    raw_citation=f"cite {i}",
                    normalized_citation=f"cite {i}",
                    citation_type="FullCaseCitation",
                    context="...",
                    paragraph_index=1,
                    metadata={},
                    bluebook_normalized=False,
                    status=status,
                    confidence=90,
                    source=None,
                    source_url=None,
                    evidence="",
                )
            )
        return AuditReport(
            source_document="test.docx",
            generated_at="2026-01-01T00:00:00+00:00",
            processing_seconds=1.0,
            citations=citations,
        )

    def test_counts(self) -> None:
        report = self._make_report([
            "Verified (CourtListener)",
            "Potential Hallucination",
            "Needs Review",
            "Skipped (Non-Case Citation)",
        ])
        assert report.total_citations == 4
        assert report.verified_count == 1
        assert report.hallucination_count == 1
        assert report.needs_review_count == 1
        assert report.skipped_count == 1

    def test_json_output(self) -> None:
        report = self._make_report(["Verified (CourtListener)"])
        output = report.to_string(format="json")
        data = json.loads(output)
        assert data["summary"]["verified"] == 1
        assert data["summary"]["needs_review"] == 0

    def test_markdown_output(self) -> None:
        report = self._make_report(["Needs Review"])
        output = report.to_string(format="markdown")
        assert "needs review" in output.lower()

    def test_html_output(self) -> None:
        report = self._make_report(["Verified (CourtListener)"])
        output = report.to_string(format="html")
        assert "<html" in output
        assert "Verified" in output


# ---------------------------------------------------------------------------
# Unit tests: Broad query building
# ---------------------------------------------------------------------------
# Unit tests: Fuzzy name matching
# ---------------------------------------------------------------------------

class TestFuzzyNameMatching:
    def test_extract_name_tokens_filters_noise(self) -> None:
        tokens = _extract_name_tokens("Smith v. Jones, Inc.")
        assert "smith" in tokens
        assert "jones" in tokens
        assert "v" not in tokens
        assert "inc" not in tokens

    def test_extract_name_tokens_handles_possessives(self) -> None:
        tokens = _extract_name_tokens("O'Brien's Estate v. City of Portland")
        assert "o'brien" in tokens or "obrien" in tokens
        assert "portland" in tokens
        assert "city" not in tokens
        assert "of" not in tokens

    def test_token_overlap_exact_match(self) -> None:
        ratio, count = _name_token_overlap("Smith v. Jones", "Smith v. Jones")
        assert ratio == 1.0
        assert count >= 2

    def test_token_overlap_abbreviation_match(self) -> None:
        """Abbreviated vs full entity name should partially match."""
        ratio, count = _name_token_overlap(
            "Cornerstone Therapeutics v. Derivative Action Litig.",
            "Cornerstone Therapeutics Inc. v. Professional Derivative Action Litigation",
        )
        assert ratio >= 0.5
        assert count >= 2

    def test_token_overlap_no_match(self) -> None:
        ratio, count = _name_token_overlap("Smith v. Jones", "Baker v. Cook")
        assert ratio == 0.0
        assert count == 0

    def test_token_overlap_partial_match(self) -> None:
        """Party name variations that share some tokens."""
        ratio, count = _name_token_overlap(
            "Carpenter v. United States",
            "Timothy Carpenter v. United States of America",
        )
        assert count >= 1  # "carpenter" should match
        assert ratio > 0

    def test_verified_with_fuzzy_match(self) -> None:
        """Fuzzy token matching should verify a case with abbreviated names."""
        checker = make_checker()
        citation = make_citation(
            raw="73 A.3d 697",
            normalized="73 A.3d 697",
            metadata={
                "year": "2013",
                "plaintiff": "Cornerstone Therapeutics",
                "defendant": "Derivative Action Litig.",
            },
        )

        def mock_http_get(url: str, **kwargs: Any) -> Tuple[Dict[str, Any], str, None]:
            return {
                "count": 1,
                "results": [{
                    "caseName": "Cornerstone Therapeutics Inc. v. Prof'l Derivative Action Litigation",
                    "case_name": "Cornerstone Therapeutics Inc. v. Prof'l Derivative Action Litigation",
                    "dateFiled": "2013-07-08",
                    "absolute_url": "/opinion/789/cornerstone/",
                    "cluster_id": 789,
                    "citation": ["73 A.3d 697"],
                }],
            }, "http://example.com", None

        checker._verifier._http_get_json = mock_http_get  # type: ignore[assignment]
        decision = checker._verify_citation(citation)
        assert decision.status.startswith("Verified"), f"Expected Verified, got: {decision.status}"

    def test_docket_search_verifies_case(self) -> None:
        """Docket search (type=d) should verify a case not in opinions DB."""
        checker = make_checker()
        citation = make_citation(
            raw="123 N.Y.S.3d 456",
            normalized="123 N.Y.S.3d 456",
            metadata={
                "year": "2021",
                "plaintiff": "Rodriguez",
                "defendant": "NYC Housing Authority",
            },
        )

        call_count = {"n": 0}

        def mock_http_get(url: str, **kwargs: Any) -> Tuple[Dict[str, Any], str, None]:
            call_count["n"] += 1
            params = kwargs.get("params", {})
            search_type = params.get("type", "")

            # Opinion searches (type=o) return nothing
            if search_type == "o":
                return {"count": 0, "results": []}, "http://example.com", None

            # Docket search (type=d) finds the case
            if search_type == "d":
                return {
                    "count": 1,
                    "results": [{
                        "caseName": "Rodriguez v. New York City Housing Authority",
                        "dateFiled": "2021-03-15",
                        "docket_id": 12345,
                        "absolute_url": "/docket/12345/rodriguez-v-nyc-housing/",
                    }],
                }, "http://example.com", None

            return {"count": 0, "results": []}, "http://example.com", None

        checker._verifier._http_get_json = mock_http_get  # type: ignore[assignment]
        decision = checker._verify_citation(citation)
        assert decision.status.startswith("Verified"), f"Expected Verified, got: {decision.status}"
        # Should have reached docket search
        assert any("docket" in a.strategy.lower() for a in decision.search_attempts)


# ---------------------------------------------------------------------------

class TestQueryBuilding:
    def test_broad_query_includes_parts(self) -> None:
        checker = make_checker()
        citation = make_citation()
        query = checker._build_broad_query(citation)
        assert "Brown" in query
        assert "347 U.S. 483" in query

    def test_case_name_from_metadata(self) -> None:
        checker = make_checker()
        name = checker._case_name_from_metadata({"plaintiff": "Brown", "defendant": "Board of Education"})
        assert name == "Brown v. Board of Education"

    def test_case_name_fallback(self) -> None:
        checker = make_checker()
        name = checker._case_name_from_metadata({"case_name": "Some Case"})
        assert name == "Some Case"


# ---------------------------------------------------------------------------
# Integration-level: process_document smoke test (mocked DOCX)
# ---------------------------------------------------------------------------

class TestProcessDocumentSmoke:
    def test_empty_document(self, tmp_path: Path) -> None:
        try:
            from docx import Document  # type: ignore
        except ImportError:
            pytest.skip("python-docx not installed")

        doc_path = tmp_path / "empty.docx"
        doc = Document()
        doc.add_paragraph("This document has no legal citations.")
        doc.save(str(doc_path))

        checker = make_checker()
        report = checker.process_document(doc_path)
        assert report.total_citations == 0

    def test_document_with_case_citation(self, tmp_path: Path) -> None:
        try:
            from docx import Document  # type: ignore
        except ImportError:
            pytest.skip("python-docx not installed")

        doc_path = tmp_path / "brief.docx"
        doc = Document()
        doc.add_paragraph(
            "The Court held in Brown v. Board of Education, 347 U.S. 483 (1954), "
            "that segregation violates the Equal Protection Clause."
        )
        doc.save(str(doc_path))

        checker = make_checker()

        def mock_http_get(url: str, **kwargs: Any) -> Tuple[Dict[str, Any], str, None]:
            return {
                "count": 1,
                "results": [{"citation": ["347 U.S. 483"], "absolute_url": "/opinion/123/"}],
            }, url, None

        checker._verifier._http_get_json = mock_http_get  # type: ignore[assignment]
        report = checker.process_document(doc_path)
        assert report.total_citations >= 1
        assert report.verified_count >= 1

    def test_document_with_statute(self, tmp_path: Path) -> None:
        try:
            from docx import Document  # type: ignore
        except ImportError:
            pytest.skip("python-docx not installed")

        doc_path = tmp_path / "statute.docx"
        doc = Document()
        doc.add_paragraph("Pursuant to 42 U.S.C. § 1983, the plaintiff seeks relief.")
        doc.save(str(doc_path))

        checker = make_checker()
        report = checker.process_document(doc_path)
        assert report.hallucination_count == 0

    def test_pdf_document(self, tmp_path: Path) -> None:
        """Process a PDF with a citation (mocked verification)."""
        try:
            import pdfplumber  # type: ignore
            from reportlab.lib.pagesizes import letter  # type: ignore
            from reportlab.pdfgen import canvas  # type: ignore
        except ImportError:
            pytest.skip("pdfplumber or reportlab not installed")

        pdf_path = tmp_path / "brief.pdf"
        c = canvas.Canvas(str(pdf_path), pagesize=letter)
        c.drawString(72, 700, "The Court held in Brown v. Board of Education, 347 U.S. 483 (1954).")
        c.save()

        checker = make_checker()

        def mock_http_get(url: str, **kwargs: Any) -> Tuple[Dict[str, Any], str, None]:
            return {
                "count": 1,
                "results": [{"citation": ["347 U.S. 483"], "absolute_url": "/opinion/123/"}],
            }, url, None

        checker._verifier._http_get_json = mock_http_get  # type: ignore[assignment]
        report = checker.process_document(pdf_path)
        assert report.total_citations >= 1

    def test_executive_summary_in_markdown(self) -> None:
        """Markdown report should contain Action Required section when hallucinations exist."""
        from legal_citation_checker.pipeline import AuditReport, CitationAudit

        report = AuditReport(
            source_document="test.docx",
            generated_at="2026-01-01T00:00:00+00:00",
            processing_seconds=1.0,
            citations=[
                CitationAudit(
                    index=1,
                    raw_citation="999 F.3d 999",
                    normalized_citation="999 F.3d 999",
                    citation_type="FullCaseCitation",
                    context="See Fake v. Case, 999 F.3d 999 (9th Cir. 2099).",
                    paragraph_index=1,
                    metadata={},
                    bluebook_normalized=False,
                    status="Potential Hallucination",
                    confidence=92,
                    source=None,
                    source_url=None,
                    evidence="Not found",
                ),
            ],
        )
        md = report.to_string(format="markdown")
        assert "Action Required" in md
        assert "999 F.3d 999" in md

    def test_document_with_westlaw_citation(self, tmp_path: Path) -> None:
        try:
            from docx import Document  # type: ignore
        except ImportError:
            pytest.skip("python-docx not installed")

        doc_path = tmp_path / "westlaw.docx"
        doc = Document()
        doc.add_paragraph(
            "See Smith v. Jones, 2025 WL 2192378 (Cal. Ct. App. 2025)."
        )
        doc.save(str(doc_path))

        checker = make_checker()
        report = checker.process_document(doc_path)
        # WL citations should be "Needs Review", not hallucination.
        assert report.hallucination_count == 0
