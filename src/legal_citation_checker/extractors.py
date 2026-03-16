"""Document text extraction and citation extraction via eyecite."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from .models import DocumentText, ExtractedCitation, ParsedCitation, ParagraphSpan
from .normalizer import CitationNormalizer

# Citation types that should be skipped (not verifiable against case law databases).
SKIP_CITATION_TYPES = frozenset({
    "NonopinionCitation",
    "IdCitation",
    "SupraCitation",
    "ShortCaseCitation",
    "UnknownCitation",
})

# Regex patterns for statutes, regulations, and other non-case citations.
_STATUTE_PATTERNS = [
    re.compile(r"\b\d+\s+U\.?S\.?C\.?\s*(§|[Ss]ection)?\s*\d+", re.IGNORECASE),
    re.compile(r"\b\d+\s+C\.?F\.?R\.?\s*(§|[Ss]ection|[Pp]art)?\s*\d+", re.IGNORECASE),
    re.compile(r"\bPub\.?\s*L\.?\s*No\.?\s*\d+", re.IGNORECASE),
    re.compile(r"\b[A-Z][a-z]+\.?\s*(Rev\.?\s*)?((Ann\.?\s*)?(Code|Stat|Laws|Gen\.?\s*Stat))", re.IGNORECASE),
]


def extract_docx_text(path: Path) -> DocumentText:
    """Extract text from a DOCX file preserving paragraph boundaries."""
    Document = _import_python_docx_document()
    doc = Document(str(path))

    paragraphs: List[ParagraphSpan] = []
    chunks: List[str] = []
    cursor = 0

    for paragraph_index, paragraph in enumerate(doc.paragraphs, start=1):
        text = (paragraph.text or "").strip()
        if not text:
            continue

        if chunks:
            chunks.append("\n\n")
            cursor += 2

        start = cursor
        chunks.append(text)
        cursor += len(text)
        end = cursor

        paragraphs.append(
            ParagraphSpan(
                index=paragraph_index,
                text=text,
                start=start,
                end=end,
            )
        )

    full_text = "".join(chunks)
    return DocumentText(full_text=full_text, paragraphs=paragraphs)


def extract_pdf_text(path: Path) -> DocumentText:
    """Extract text from a PDF file using pdfplumber or PyPDF2."""
    try:
        import pdfplumber  # type: ignore
    except ImportError:
        try:
            from PyPDF2 import PdfReader  # type: ignore
        except ImportError:
            raise RuntimeError(
                "PDF support requires pdfplumber or PyPDF2. "
                "Install with: pip install pdfplumber"
            )
        return _extract_pdf_text_pypdf2(path, PdfReader)
    return _extract_pdf_text_pdfplumber(path, pdfplumber)


def _extract_pdf_text_pdfplumber(path: Path, pdfplumber: Any) -> DocumentText:
    paragraphs: List[ParagraphSpan] = []
    chunks: List[str] = []
    cursor = 0

    with pdfplumber.open(str(path)) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            text = (page.extract_text() or "").strip()
            if not text:
                continue

            if chunks:
                chunks.append("\n\n")
                cursor += 2

            start = cursor
            chunks.append(text)
            cursor += len(text)

            paragraphs.append(
                ParagraphSpan(index=page_num, text=text, start=start, end=cursor)
            )

    return DocumentText(full_text="".join(chunks), paragraphs=paragraphs)


def _extract_pdf_text_pypdf2(path: Path, PdfReader: Any) -> DocumentText:
    paragraphs: List[ParagraphSpan] = []
    chunks: List[str] = []
    cursor = 0

    reader = PdfReader(str(path))
    for page_num, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            continue

        if chunks:
            chunks.append("\n\n")
            cursor += 2

        start = cursor
        chunks.append(text)
        cursor += len(text)

        paragraphs.append(
            ParagraphSpan(index=page_num, text=text, start=start, end=cursor)
        )

    return DocumentText(full_text="".join(chunks), paragraphs=paragraphs)


def extract_citations(
    document_text: DocumentText,
    normalizer: CitationNormalizer,
    verbose: bool = False,
) -> List[ExtractedCitation]:
    """Extract and normalize citations from document text using eyecite."""
    get_citations = _import_eyecite_get_citations()

    if not document_text.full_text.strip():
        return []

    citations: List[ExtractedCitation] = []
    extracted = get_citations(document_text.full_text)

    for i, citation_obj in enumerate(extracted, start=1):
        citation_type_name = type(citation_obj).__name__

        if citation_type_name in SKIP_CITATION_TYPES:
            continue

        span = _citation_span(citation_obj)
        raw_citation = _raw_citation_text(citation_obj, document_text.full_text, span)

        # Skip bare symbols and very short non-citation text.
        stripped = raw_citation.strip(" §.,;:()")
        if len(stripped) < 3:
            continue

        # Skip statute/regulation patterns.
        if is_statute_or_regulation(raw_citation):
            continue

        metadata = _citation_metadata(citation_obj)
        parsed = _build_parsed_citation(citation_obj, metadata)
        normalized, changed = normalizer.normalize(raw_citation, metadata)
        paragraph_index = _paragraph_for_span(document_text.paragraphs, span)
        context = _citation_context(document_text, span)

        citations.append(
            ExtractedCitation(
                index=len(citations) + 1,
                raw_citation=raw_citation,
                normalized_citation=normalized,
                citation_type=citation_type_name,
                context=context,
                paragraph_index=paragraph_index,
                metadata=metadata,
                bluebook_normalized=changed,
                parsed=parsed,
            )
        )

    # Fallback: catch citations eyecite missed (e.g. unrecognized reporter editions)
    _fallback_extract_missed_citations(
        document_text, normalizer, citations, verbose,
    )

    return citations


# Regex for common citation formats that eyecite may not recognise
# (e.g. future reporter editions like Cal.App.6th, Cal.App.7th, etc.)
_FALLBACK_CITE_RE = re.compile(
    r"(?P<volume>\d{1,4})\s+"
    r"(?P<reporter>"
    r"Cal\.?\s*App\.?\s*(?:6th|7th|8th|9th|10th)"  # future CA App editions
    r"|Cal\.?\s*(?:6th|7th|8th)"                     # future CA editions
    r")"
    r"\s+(?P<page>\d{1,5})"
)


def _fallback_extract_missed_citations(
    document_text: DocumentText,
    normalizer: CitationNormalizer,
    existing: List[ExtractedCitation],
    verbose: bool,
) -> None:
    """Regex fallback for citations with reporter editions eyecite doesn't know."""
    # Build set of already-extracted raw spans to avoid duplicates
    existing_spans = set()
    for c in existing:
        existing_spans.add(c.raw_citation.strip())

    text = document_text.full_text
    for m in _FALLBACK_CITE_RE.finditer(text):
        raw = m.group(0).strip()
        if raw in existing_spans:
            continue

        volume = m.group("volume")
        reporter = m.group("reporter")
        page = m.group("page")

        # Try to find party names from surrounding context
        prefix = text[max(0, m.start() - 100):m.start()]
        plaintiff, defendant = _extract_parties_from_prefix(prefix)

        # Try to find year from suffix
        suffix = text[m.end():m.end() + 30]
        year_match = re.search(r"\((\d{4})\)", suffix)
        year = year_match.group(1) if year_match else None

        metadata: Dict[str, Any] = {
            "volume": volume, "reporter": reporter, "page": page,
            "plaintiff": plaintiff, "defendant": defendant, "year": year,
        }
        parsed = ParsedCitation(
            volume=volume, reporter=reporter, page=page,
            plaintiff=plaintiff, defendant=defendant, year=year,
        )
        normalized, changed = normalizer.normalize(raw, metadata)
        span = (m.start(), m.end())
        paragraph_index = _paragraph_for_span(document_text.paragraphs, span)
        context = _citation_context(document_text, span)

        existing.append(
            ExtractedCitation(
                index=len(existing) + 1,
                raw_citation=raw,
                normalized_citation=normalized,
                citation_type="FullCaseCitation",
                context=context,
                paragraph_index=paragraph_index,
                metadata=metadata,
                bluebook_normalized=changed,
                parsed=parsed,
            )
        )
        existing_spans.add(raw)


def _extract_parties_from_prefix(prefix: str) -> Tuple[Optional[str], Optional[str]]:
    """Try to extract party names from text immediately before a citation."""
    # Look for "Party1 v. Party2, " pattern
    m = re.search(
        r"([A-Z][A-Za-z'\-.\s]+?)\s+v\.\s+([A-Z][A-Za-z'\-.\s]+?),\s*$",
        prefix,
    )
    if m:
        return m.group(1).strip().strip("*"), m.group(2).strip().strip("*")
    return None, None


def is_statute_or_regulation(citation_text: str) -> bool:
    """Check if citation text matches a statute or regulation pattern."""
    return any(pattern.search(citation_text) for pattern in _STATUTE_PATTERNS)


def _build_parsed_citation(citation_obj: Any, metadata: Dict[str, Any]) -> ParsedCitation:
    """Extract structured citation components from eyecite object + metadata."""
    # eyecite stores volume/reporter/page in groups dict
    groups = getattr(citation_obj, "groups", {}) or {}
    volume = groups.get("volume")
    reporter = groups.get("reporter")
    page = groups.get("page")

    # Reporter full name from edition info
    reporter_full_name = None
    cite_type = None
    edition = getattr(citation_obj, "edition_guess", None)
    if edition is not None:
        rep_obj = getattr(edition, "reporter", None)
        if rep_obj is not None:
            reporter_full_name = getattr(rep_obj, "name", None)
            cite_type = getattr(rep_obj, "cite_type", None)

    # Party names and court/year from metadata (already extracted)
    # Strip markdown italic markers (*) that may leak through from .md/.docx text
    plaintiff = str(metadata.get("plaintiff") or "").strip().strip("*").strip() or None
    defendant = str(metadata.get("defendant") or "").strip().strip("*").strip() or None
    year = str(metadata.get("year") or "").strip() or None
    court = str(metadata.get("court") or "").strip() or None

    # Pincite from eyecite metadata
    eyecite_meta = getattr(citation_obj, "metadata", None)
    pincite = None
    if eyecite_meta is not None:
        pincite = getattr(eyecite_meta, "pin_cite", None)

    return ParsedCitation(
        volume=volume,
        reporter=reporter,
        page=page,
        plaintiff=plaintiff,
        defendant=defendant,
        year=year,
        court=court,
        pincite=pincite,
        reporter_full_name=reporter_full_name,
        cite_type=cite_type,
    )


def _citation_span(citation_obj: Any) -> Optional[Tuple[int, int]]:
    for attr in ("full_span", "span"):
        value = getattr(citation_obj, attr, None)
        if callable(value):
            try:
                value = value()
            except Exception:
                value = None
        if (
            isinstance(value, tuple)
            and len(value) == 2
            and all(isinstance(v, int) for v in value)
            and value[0] >= 0
            and value[1] >= value[0]
        ):
            return value
    return None


def _raw_citation_text(
    citation_obj: Any,
    full_text: str,
    span: Optional[Tuple[int, int]],
) -> str:
    for attr in ("corrected_citation", "matched_text", "text"):
        value = getattr(citation_obj, attr, None)
        if callable(value):
            try:
                value = value()
            except Exception:
                value = None
        if isinstance(value, str) and value.strip():
            return value.strip()

    if span is not None:
        start, end = span
        snippet = full_text[start:end].strip()
        if snippet:
            return snippet

    fallback = str(citation_obj).strip()
    return fallback if fallback else "<unknown citation>"


def _citation_metadata(citation_obj: Any) -> Dict[str, Any]:
    metadata: Dict[str, Any] = {}

    raw_meta = getattr(citation_obj, "metadata", None)
    if raw_meta is not None:
        if isinstance(raw_meta, dict):
            for key, value in raw_meta.items():
                _set_serializable(metadata, str(key), value)
        else:
            for key in dir(raw_meta):
                if key.startswith("_"):
                    continue
                try:
                    value = getattr(raw_meta, key)
                except Exception:
                    continue
                if callable(value):
                    continue
                _set_serializable(metadata, key, value)

    for attr in (
        "year",
        "court",
        "plaintiff",
        "defendant",
        "volume",
        "reporter",
        "page",
    ):
        if attr in metadata:
            continue
        try:
            value = getattr(citation_obj, attr)
        except Exception:
            value = None
        _set_serializable(metadata, attr, value)

    return metadata


def _set_serializable(target: Dict[str, Any], key: str, value: Any) -> None:
    if value is None:
        return
    if isinstance(value, (str, int, float, bool)):
        target[key] = value
        return
    if isinstance(value, list):
        serializable = [item for item in value if isinstance(item, (str, int, float, bool))]
        if serializable:
            target[key] = serializable
        return
    target[key] = str(value)


def _paragraph_for_span(
    paragraphs: List[ParagraphSpan],
    span: Optional[Tuple[int, int]],
) -> Optional[int]:
    if span is None:
        return None
    start, _ = span
    for paragraph in paragraphs:
        if paragraph.start <= start < paragraph.end:
            return paragraph.index
    return None


def _citation_context(
    document_text: DocumentText,
    span: Optional[Tuple[int, int]],
    radius: int = 140,
) -> str:
    if not document_text.full_text:
        return ""

    if span is None:
        first_para = document_text.paragraphs[0].text if document_text.paragraphs else ""
        return first_para[: radius * 2]

    start, end = span
    left = max(0, start - radius)
    right = min(len(document_text.full_text), end + radius)
    snippet = document_text.full_text[left:right]
    return re.sub(r"\s+", " ", snippet).strip()


def _import_python_docx_document() -> Any:
    try:
        from docx import Document  # type: ignore
    except Exception as exc:
        raise RuntimeError(
            "python-docx is required for DOCX processing. "
            "Install dependencies with: pip install -r requirements.txt"
        ) from exc
    return Document


def _import_eyecite_get_citations() -> Any:
    try:
        from eyecite import get_citations  # type: ignore
    except Exception as exc:
        raise RuntimeError(
            "eyecite is required for citation extraction. "
            "Install dependencies with: pip install -r requirements.txt"
        ) from exc
    return get_citations
